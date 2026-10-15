"""資料ファイルからの本文抽出。位置情報(ページ・段落)を保持する。

対応: PDF(テキスト / OCR)、Word(.docx)、Excel(.xlsx)、メール(.eml / .msg)、テキスト(.txt .md .csv .html)
任意の依存: PyMuPDF(PDF)、pytesseract + Tesseract(OCR)、python-docx、openpyxl、extract-msg
"""
from __future__ import annotations

import email
import email.policy
import html
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".log"}
HTML_EXT = {".html", ".htm"}
SUPPORTED = TEXT_EXT | HTML_EXT | {".pdf", ".docx", ".xlsx", ".eml", ".msg"}
MAX_BYTES = 200 * 1024 * 1024
MIN_PAGE_CHARS = 20   # これ未満のPDFページはスキャン画像とみなしOCRを試みる


@dataclass
class Part:
    loc: str              # 表示用の位置(p.3 / 段落1〜12 / 本文)
    text: str
    page: int | None = None


@dataclass
class Extracted:
    parts: list[Part]
    method: str = "text"  # text / ocr / mixed
    pages: int = 0
    message: str = ""


class ExtractError(Exception):
    pass


def clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = re.sub(r"[ \t　]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def extract(path: Path, ocr_lang: str = "jpn") -> Extracted:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractError("未対応の形式")
    if path.stat().st_size > MAX_BYTES:
        raise ExtractError("ファイルが大きすぎます")
    if ext in TEXT_EXT:
        return Extracted([Part("本文", clean(_read_text(path)))])
    if ext in HTML_EXT:
        return Extracted([Part("本文", clean(_strip_html(_read_text(path))))])
    if ext == ".pdf":
        return _pdf(path, ocr_lang)
    if ext == ".docx":
        return _docx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    if ext == ".eml":
        return _eml(path)
    if ext == ".msg":
        return _msg(path)
    raise ExtractError("未対応の形式")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "euc_jp", "iso2022_jp"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _strip_html(s: str) -> str:
    s = re.sub(r"(?is)<(script|style).*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", s)
    return html.unescape(re.sub(r"<[^>]+>", " ", s))


# ---------------------------------------------------------------- PDF

def _ocr_available() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001
        return False


def _pdf(path: Path, ocr_lang: str) -> Extracted:
    try:
        import pymupdf
    except ImportError as e:
        raise ExtractError("PyMuPDF が未インストールです(pip install pymupdf)") from e
    try:
        doc = pymupdf.open(path)
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"PDFを開けません: {e}") from e
    if doc.needs_pass:
        raise ExtractError("パスワード付きPDF")
    parts, ocr_pages, no_text = [], 0, 0
    ocr_ok = None
    for page in doc:
        n = page.number + 1
        text = clean(page.get_text())
        if len(text) < MIN_PAGE_CHARS:
            if ocr_ok is None:
                ocr_ok = _ocr_available()
            if ocr_ok:
                text = clean(_ocr_page(page, ocr_lang)) or text
                ocr_pages += 1
            elif not text:
                no_text += 1
        if text:
            parts.append(Part(f"p.{n}", text, n))
    method = "ocr" if ocr_pages and ocr_pages == len(doc) else "mixed" if ocr_pages else "text"
    msg = ""
    if no_text:
        msg = f"文字のないページ {no_text} 件(スキャン画像の可能性。OCR未設定)"
    return Extracted(parts, method, len(doc), msg)


def _ocr_page(page, lang: str) -> str:
    import io

    import pytesseract
    from PIL import Image

    pix = page.get_pixmap(dpi=300)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    return pytesseract.image_to_string(img, lang=lang)


# ---------------------------------------------------------------- Office

def _docx(path: Path) -> Extracted:
    try:
        import docx
    except ImportError as e:
        raise ExtractError("python-docx が未インストールです") from e
    d = docx.Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text for c in row.cells))
    # 段落番号つきで 15 段落ずつに区切る(出典に「段落n〜m」を表示する)
    parts, buf, start = [], [], 1
    for i, p in enumerate(paras, 1):
        if p.strip():
            buf.append(p)
        if len(buf) >= 15 or (i == len(paras) and buf):
            parts.append(Part(f"段落{start}〜{i}", clean("\n".join(buf))))
            buf, start = [], i + 1
    return Extracted(parts, "text", len(paras))


def _xlsx(path: Path) -> Extracted:
    try:
        import openpyxl
    except ImportError as e:
        raise ExtractError("openpyxl が未インストールです") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                rows.append(" | ".join(vals))
        if rows:
            parts.append(Part(f"シート「{ws.title}」", clean("\n".join(rows))))
    return Extracted(parts, "text", len(wb.worksheets))


# ---------------------------------------------------------------- メール

def _eml(path: Path) -> Extracted:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = "\n".join(f"{k}: {msg.get(k, '')}" for k in ("Date", "From", "To", "Cc", "Subject") if msg.get(k))
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = _strip_html(text)
    atts = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    if atts:
        text += "\n\n[添付] " + "、".join(atts)
    return Extracted([Part("メール", clean(head + "\n\n" + text))])


def _msg(path: Path) -> Extracted:
    try:
        import extract_msg
    except ImportError as e:
        raise ExtractError("Outlook .msg には extract-msg が必要です(pip install extract-msg)") from e
    m = extract_msg.Message(str(path))
    head = f"Date: {m.date}\nFrom: {m.sender}\nTo: {m.to}\nSubject: {m.subject}"
    return Extracted([Part("メール", clean(head + "\n\n" + (m.body or "")))])
