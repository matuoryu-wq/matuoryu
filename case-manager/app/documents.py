"""書面ひな形への差し込み、引継ぎ書・報告文の Word 出力。"""
from __future__ import annotations

import datetime as dt
import io
import re
from pathlib import Path

from . import db

PLACEHOLDER = re.compile(r"\{\{\s*([^{}\s]+)\s*\}\}")
WEEKDAYS = "月火水木金土日"

OFFICE_KEYS = {
    "事務所名": "office_name", "弁護士名": "lawyer_name", "事務所住所": "office_address",
    "事務所電話": "office_phone", "事務所FAX": "office_fax", "所属弁護士会": "bar_association",
    "登録番号": "registration_no",
}


# ---------------------------------------------------------------- 日付

def wareki(d: dt.date) -> str:
    """和暦(令和・平成)表記。"""
    if d >= dt.date(2019, 5, 1):
        era, y = "令和", d.year - 2018
    elif d >= dt.date(1989, 1, 8):
        era, y = "平成", d.year - 1988
    else:
        return f"{d.year}年{d.month}月{d.day}日"
    return f"{era}{'元' if y == 1 else y}年{d.month}月{d.day}日"


def jdate(iso: str | None, era: bool = True) -> str:
    if not iso:
        return ""
    d = dt.date.fromisoformat(iso[:10])
    return wareki(d) if era else f"{d.year}年{d.month}月{d.day}日"


# ---------------------------------------------------------------- 差し込みデータ

def office_info(conn) -> dict[str, str]:
    return {label: db.get_setting(conn, key, "") or "" for label, key in OFFICE_KEYS.items()}


def merge_fields(conn, case_id: int) -> dict[str, str]:
    c = dict(conn.execute("SELECT * FROM cases WHERE id=?", (case_id,)).fetchone())
    parties = [dict(r) for r in conn.execute("SELECT * FROM parties WHERE case_id=? ORDER BY id", (case_id,))]

    def by_role(*roles):
        return [p for p in parties if p["role"] in roles]

    clients = by_role("依頼者") or by_role("相談者")
    adverse = by_role("相手方")
    agents = by_role("相手方代理人")
    nxt = conn.execute("SELECT * FROM deadlines WHERE case_id=? AND done=0 AND kind='期日' AND due_date>=? "
                       "ORDER BY due_date, due_time LIMIT 1", (case_id, db.today().isoformat())).fetchone()
    today = db.today()
    f = {
        "事件名": c["title"], "事件番号": c["case_number"] or "", "裁判所": c["court"] or "", "分野": c["field"],
        "受任日": jdate(c["accepted_on"]), "報酬条件": c["fee_terms"] or "",
        "依頼者": "、".join(p["name"] for p in clients), "依頼者フリガナ": "、".join(p["kana"] or "" for p in clients),
        "依頼者住所": clients[0]["address"] if clients else "", "依頼者電話": clients[0]["phone"] if clients else "",
        "相手方": "、".join(p["name"] for p in adverse), "相手方住所": adverse[0]["address"] if adverse else "",
        "相手方代理人": "、".join(p["name"] for p in agents),
        "今日": wareki(today), "今日西暦": f"{today.year}年{today.month}月{today.day}日",
        "次回期日": (f"{jdate(nxt['due_date'])}({WEEKDAYS[dt.date.fromisoformat(nxt['due_date']).weekday()]})"
                 f"{' ' + nxt['due_time'] if nxt['due_time'] else ''}") if nxt else "",
        "次回期日内容": nxt["title"] if nxt else "",
        "次回期日場所": (nxt["place"] or "") if nxt else "",
    }
    f.update(office_info(conn))
    return f


FIELD_HELP = ["事件名", "事件番号", "裁判所", "分野", "受任日", "報酬条件", "依頼者", "依頼者フリガナ", "依頼者住所",
              "依頼者電話", "相手方", "相手方住所", "相手方代理人", "今日", "今日西暦", "次回期日", "次回期日内容",
              "次回期日場所", *OFFICE_KEYS]


def fill_text(text: str, fields: dict[str, str], missing: set[str]) -> str:
    def repl(m: re.Match) -> str:
        k = m.group(1)
        if k in fields and fields[k]:
            return fields[k]
        missing.add(k)
        return f"〔{k}〕"   # 未入力の項目は目立つ形で残す
    return PLACEHOLDER.sub(repl, text)


# ---------------------------------------------------------------- Word

def _paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def fill_docx(src: Path | bytes, fields: dict[str, str]) -> tuple[bytes, list[str]]:
    """{{項目}} を差し込む。段落内で書式が分かれていても置換できるよう、段落単位で連結して置換する。"""
    import docx

    doc = docx.Document(io.BytesIO(src) if isinstance(src, bytes) else str(src))
    missing: set[str] = set()
    for p in _paragraphs(doc):
        full = "".join(r.text for r in p.runs)
        if "{{" not in full:
            continue
        new = fill_text(full, fields, missing)
        if new != full and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), sorted(missing)


def placeholders_in_docx(data: bytes) -> list[str]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    found = []
    for p in _paragraphs(doc):
        found += PLACEHOLDER.findall("".join(r.text for r in p.runs))
    return sorted(set(found))


def _new_doc():
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)   # A4
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)
    st = d.styles["Normal"]
    st.font.name = "MS 明朝"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "MS 明朝")
    return d


def text_to_docx(text: str) -> bytes:
    """組込みひな形(テキスト)を Word にする。行頭「# 」は表題(中央・太字)、「>」は右寄せ。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = _new_doc()
    for line in text.splitlines():
        if line.startswith("# "):
            p = d.add_paragraph()
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(16)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith(">"):
            p = d.add_paragraph(line[1:].strip())
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            d.add_paragraph(line)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- 組込みひな形

BUILTIN_DIR = Path(__file__).parent / "doc_templates"


def builtin_templates() -> list[dict]:
    out = []
    for p in sorted(BUILTIN_DIR.glob("*.txt")):
        first = p.read_text(encoding="utf-8").splitlines()[0]
        out.append({"id": f"builtin:{p.stem}", "name": first.removeprefix("# ").strip() or p.stem,
                    "builtin": True, "file": p.name})
    return out


def render_builtin(stem: str, fields: dict[str, str]) -> tuple[bytes, list[str], str]:
    path = BUILTIN_DIR / f"{stem}.txt"
    if not path.is_file() or path.parent != BUILTIN_DIR:
        raise FileNotFoundError(stem)
    missing: set[str] = set()
    text = fill_text(path.read_text(encoding="utf-8"), fields, missing)
    return text_to_docx(text), sorted(missing), text


# ---------------------------------------------------------------- 引継ぎ書

def handover_docx(cases: list[dict], office: dict[str, str], note: str = "") -> bytes:
    """第三者の弁護士に渡す案件一覧(各案件の現況)。cases は case_snapshot() の結果。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = _new_doc()
    p = d.add_paragraph()
    r = p.add_run("事件引継書")
    r.bold, r.font.size = True, Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(f"作成日 {wareki(db.today())}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if office.get("事務所名") or office.get("弁護士名"):
        d.add_paragraph(f"{office.get('事務所名', '')} {office.get('弁護士名', '')}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("取扱注意: 本書には依頼者の秘密が含まれます。受領者以外への開示・複製をしないでください。")
    if note:
        d.add_paragraph(note)

    d.add_heading(f"案件一覧({len(cases)}件)", level=1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["No.", "事件名", "状態", "事件番号・裁判所", "直近の期限・期日"]):
        t.rows[0].cells[i].text = h
    for i, c in enumerate(cases, 1):
        nd = c["deadlines"][0] if c["deadlines"] else None
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = str(i), c["title"], c["status"]
        row[3].text = f"{c['case_number']} {c['court']}".strip()
        row[4].text = f"{jdate(nd['due_date'])} {nd['title']}" if nd else "—"

    for i, c in enumerate(cases, 1):
        d.add_page_break()
        d.add_heading(f"{i}. {c['title']}", level=1)
        info = [("状態", c["status"]), ("分野", c["field"]), ("受任日", jdate(c["accepted_on"])),
                ("裁判所", c["court"]), ("事件番号", c["case_number"]), ("受任経路", c.get("referral_source") or ""),
                ("報酬条件", c["fee_terms"]), ("記録フォルダ", c["folder_path"])]
        tb = d.add_table(rows=0, cols=2)
        tb.style = "Table Grid"
        for k, v in info:
            cells = tb.add_row().cells
            cells[0].text, cells[1].text = k, v or "—"

        d.add_heading("当事者", level=2)
        for p_ in c["parties"]:
            contact = " / ".join(x for x in (p_["phone"], p_["email"], p_["address"]) if x)
            d.add_paragraph(f"{p_['role']}: {p_['name']}{'(' + p_['kana'] + ')' if p_['kana'] else ''}"
                            f"{'  ' + contact if contact else ''}", style="List Bullet")
        if not c["parties"]:
            d.add_paragraph("登録なし")

        d.add_heading("未了の期限・期日", level=2)
        for x in c["deadlines"]:
            flag = "" if x["confirmed"] else "【計算値・未確定】"
            d.add_paragraph(f"{jdate(x['due_date'])} {x['due_time'] or ''} {x['kind']} {x['title']} {x['place'] or ''}{flag}",
                            style="List Bullet")
        if not c["deadlines"]:
            d.add_paragraph("なし")

        d.add_heading("対応中の事項(TODO・待ち)", level=2)
        for x in c["todos"]:
            wait = f"(待ち: {x['waiting_for']})" if x["waiting_for"] else ""
            due = f" 期限 {jdate(x['due_date'])}" if x["due_date"] else ""
            d.add_paragraph(f"[{x['status']}] {x['title']}{wait}{due}", style="List Bullet")
        if not c["todos"]:
            d.add_paragraph("なし")

        d.add_heading("会計", level=2)
        fs = c["finance"]
        d.add_paragraph(f"報酬: 請求 {fs['fees']['billed']:,}円 / 入金 {fs['fees']['received']:,}円 / 未収 {fs['fees']['receivable']:,}円")
        d.add_paragraph(f"実費: 立替 {fs['expenses']['spent']:,}円 / 回収 {fs['expenses']['recovered']:,}円")
        d.add_paragraph(f"預り金残高: {fs['deposit']['balance']:,}円(受入 {fs['deposit']['received']:,}円)")

        d.add_heading("最近の経過(新しい順)", level=2)
        for a in c["activities"]:
            first = a["content"].splitlines()[0] if a["content"] else ""
            d.add_paragraph(f"{jdate(a['date'])} [{a['kind']}] {first[:120]}", style="List Bullet")
        if not c["activities"]:
            d.add_paragraph("記録なし")
        if c["memo"]:
            d.add_heading("メモ・引継ぎ上の注意", level=2)
            for line in c["memo"].splitlines():
                d.add_paragraph(line)
        if c["checklist_open"]:
            d.add_heading("受任時チェックリストの未了項目", level=2)
            for x in c["checklist_open"]:
                d.add_paragraph(x, style="List Bullet")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def letter_docx(to: str, subject: str, paragraphs: list[str], office: dict[str, str]) -> bytes:
    """依頼者への報告文(下書き)を手紙形式の Word に。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = _new_doc()
    d.add_paragraph(wareki(db.today())).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph(f"{to} 様")
    for k in ("事務所名", "事務所住所", "事務所電話"):
        if office.get(k):
            d.add_paragraph(office[k]).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if office.get("弁護士名"):
        d.add_paragraph(f"弁護士 {office['弁護士名']}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = d.add_paragraph(subject)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for para in paragraphs:
        d.add_paragraph(para)
    d.add_paragraph("以上").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
