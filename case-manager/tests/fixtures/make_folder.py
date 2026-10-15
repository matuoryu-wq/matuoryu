"""模擬事件フォルダ(髙橋事件)を作る。PDF・Word は生成、メール・テキストは fixtures からコピー。"""
import shutil
from pathlib import Path

SRC = Path(__file__).parent / "folder_takahashi"


def build(dest: Path) -> Path:
    import docx
    import pymupdf

    dest.mkdir(parents=True, exist_ok=True)
    for f in SRC.iterdir():
        shutil.copy(f, dest / f.name)

    doc = pymupdf.open()
    p1 = doc.new_page()
    y = 72
    for line in ["売買契約書", "売主 髙橋一郎(以下「甲」という。)", "買主 株式会社サンプル商事(以下「乙」という。)",
                 "第1条 甲は乙に対し、機械部品一式を代金320万円で売り渡す。"]:
        p1.insert_text((50, y), line, fontname="japan", fontsize=11)
        y += 20
    p2 = doc.new_page()
    y = 72
    for line in ["第5条 乙は、納品後7日以内に検収を行い、異議がない場合は検収完了とみなす。",
                 "第6条 代金は2026年5月31日までに支払う。", "締結日 2026年2月1日"]:
        p2.insert_text((50, y), line, fontname="japan", fontsize=11)
        y += 20
    (dest / "契約").mkdir(exist_ok=True)
    doc.save(dest / "契約" / "売買契約書.pdf")

    scan = pymupdf.open()
    sp = scan.new_page()
    sp.draw_rect(pymupdf.Rect(50, 50, 300, 200), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))  # 文字のない画像相当
    scan.save(dest / "納品書_スキャン.pdf")

    d = docx.Document()
    d.add_paragraph("打合せメモ 2026年5月10日")
    d.add_paragraph("依頼者 髙橋一郎氏。相手方担当者は佐東氏(依頼者は「佐藤」と記憶)。")
    d.add_paragraph("相手方は寸法不具合を理由に支払を保留。検収書は未受領。")
    d.save(dest / "打合せメモ.docx")
    return dest
