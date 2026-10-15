"""v3: 書面ひな形・現状要約・報告文・カレンダー・保管期限・引継ぎ書・受任経路分析。"""
import datetime as dt
import io
import sqlite3
import urllib.parse

import docx
import pytest

from conftest import TOKEN

TAKAHASHI = "髙橋 対 ㈱サンプル商事 売買代金請求"
KEIJI = "刑事 窃盗被告事件(被告人 山﨑)"


def doc_text(content: bytes) -> str:
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def expire_reauth(settings):
    db = sqlite3.connect(settings.db_path)
    db.execute("UPDATE sessions SET reauth_at=?", ((dt.datetime.now() - dt.timedelta(hours=1)).isoformat(),))
    db.commit()
    db.close()


@pytest.fixture
def office(client):
    client.put("/api/settings/office", json={"事務所名": "テスト法律事務所", "弁護士名": "甲野 太郎",
                                             "事務所住所": "東京都千代田区霞が関1-1-1", "事務所電話": "03-0000-0000",
                                             "所属弁護士会": "第一東京弁護士会"})


# ---------------------------------------------------------------- 書面ひな形

def test_builtin_template(client, loaded, office):
    cid = loaded[TAKAHASHI]
    tpls = client.get("/api/doc-templates").json()
    assert {t["name"] for t in tpls if t["builtin"]} >= {"委任状", "受任通知書", "FAX送信書"}
    chk = client.get(f"/api/cases/{cid}/documents/check", params={"template": "builtin:01_委任状"}).json()
    assert chk["missing"] == ["事務所FAX", "依頼者住所"]
    r = client.get(f"/api/cases/{cid}/documents/render", params={"template": "builtin:01_委任状"})
    assert r.status_code == 200
    assert urllib.parse.quote("委任状_") in r.headers["content-disposition"]
    text = doc_text(r.content)
    for w in ["髙橋 一郎", "株式会社サンプル商事", "令和8年(ワ)第1234号", "東京地方裁判所", "甲野 太郎", "第一東京弁護士会",
              "復代理人の選任", "〔依頼者住所〕"]:
        assert w in text, w
    assert "{{" not in text
    assert any(a["action"] == "書面の作成" for a in client.get("/api/audit").json())


def test_uploaded_template_with_split_runs(client, loaded, office):
    cid = loaded[TAKAHASHI]
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("原告 {{依")          # 書式の違いでランが分かれていても置換できる
    p.add_run("頼者}} 被告 {{相手方}}")
    d.add_paragraph("事件番号 {{事件番号}} 次回 {{次回期日}} {{謎の項目}}")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "裁判所: {{裁判所}}"
    buf = io.BytesIO()
    d.save(buf)
    up = client.post("/api/doc-templates", files={"file": ("訴状表紙.docx", buf.getvalue(), "application/octet-stream")},
                     data={"name": "訴状表紙"}).json()
    assert up["unknown_fields"] == ["謎の項目"]
    r = client.get(f"/api/cases/{cid}/documents/render", params={"template": str(up["id"])})
    text = doc_text(r.content)
    assert "原告 髙橋 一郎 被告 株式会社サンプル商事" in text
    assert "令和8年11月5日(木) 10:30" in text
    assert "裁判所: 東京地方裁判所 民事第〇部" in text and "〔謎の項目〕" in text
    assert client.post("/api/doc-templates", files={"file": ("x.txt", b"a", "text/plain")}).status_code == 422
    assert client.delete(f"/api/doc-templates/{up['id']}").status_code == 200
    assert client.get(f"/api/cases/{cid}/documents/render", params={"template": str(up["id"])}).status_code == 404
    assert client.get(f"/api/cases/{cid}/documents/render", params={"template": "builtin:../../x"}).status_code == 404


# ---------------------------------------------------------------- 現状要約・報告文

def test_case_summary(client, loaded):
    cid = loaded[TAKAHASHI]
    s = client.get(f"/api/cases/{cid}/summary").json()
    assert s["deadlines"][0]["title"] == "第1回口頭弁論期日"
    assert s["todos"][0]["waiting_for"] == "相手方答弁書の提出待ち"
    assert s["checklist_open"] and s["finance"]["deposit"]["balance"] == 0
    assert len(s["parties"]) == 3 and s["worktime_minutes"] == 120


def test_ai_status_and_report(client, loaded, fake_llm, office):
    cid = loaded[TAKAHASHI]
    for kind in ("status", "report"):
        job = client.post(f"/api/cases/{cid}/ai/run", json={"kind": kind}).json()
        j = client.get(f"/api/jobs/{job['job_id']}").json()
        assert j["status"] == "完了", j
        run = client.get(f"/api/ai/runs/{j['result_id']}").json()
        res = run["result"]
        if kind == "status":
            assert res["headline"] and res["next_steps"][0]["citations"]
            assert client.get(f"/api/cases/{cid}/summary").json()["ai_status"]["run_id"] == run["id"]
        else:
            assert res["paragraphs"][0]["text"].startswith("髙橋 一郎様")      # 記号が復元される
            r = client.get(f"/api/ai/runs/{run['id']}/docx")
            text = doc_text(r.content)
            assert "髙橋 一郎 様" in text and "ご依頼の件の経過について" in text and "テスト法律事務所" in text
            assert "弁護士 甲野 太郎" in text
            # 期限・TODO も資料として送られている(匿名化済み)
            sent = fake_llm.calls[-1]["user"]
            assert "第1回口頭弁論期日" in sent and "相手方答弁書の提出待ち" in sent and "髙橋" not in sent
    acts = client.get(f"/api/cases/{cid}").json()["activities"]
    assert len(acts) == 2   # 報告文は自動で活動記録にしない(送信は手動)


# ---------------------------------------------------------------- カレンダー

def test_calendar_feed(app, client, loaded, settings):
    from fastapi.testclient import TestClient

    cid = loaded[TAKAHASHI]
    client.post(f"/api/cases/{cid}/deadlines", json={"title": "控訴期間", "due_date": "2026-12-01",
                                                     "calc": {"steps": [], "basis": [], "warnings": []}})
    kid = loaded[KEIJI]
    client.post(f"/api/cases/{kid}/deadlines", json={"title": "機微な期日", "due_date": "2026-12-02", "kind": "期日",
                                                     "due_time": "13:30", "place": "秘匿の法廷"})
    client.post(f"/api/cases/{cid}/todos", json={"title": "証拠整理", "due_date": "2026-11-20"})

    expire_reauth(settings)
    assert client.post("/api/calendar/token").status_code == 403
    client.post("/api/auth/reauth", json={"token": TOKEN})
    path = client.post("/api/calendar/token").json()["path"]

    anon = TestClient(app)                      # カレンダーアプリは Cookie なし
    r = anon.get(path)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/calendar")
    ics = r.text
    assert ics.startswith("BEGIN:VCALENDAR") and ics.count("BEGIN:VEVENT") == 4
    unfolded = ics.replace("\r\n ", "")
    assert "期日: 第1回口頭弁論期日(髙橋 対 ㈱サンプル商事 売買代金請求)" in unfolded
    assert "DTSTART;TZID=Asia/Tokyo:20261105T103000" in unfolded
    assert "【未確定】期限: 控訴期間" in unfolded and "DTSTART;VALUE=DATE:20261201" in unfolded
    assert f"機微な期日(案件#{kid})" in unfolded and "窃盗" not in unfolded and "秘匿の法廷" not in unfolded
    assert "TODO: 証拠整理" in unfolded and "BEGIN:VALARM" in unfolded
    assert all(len(line.encode()) <= 75 for line in ics.split("\r\n"))

    assert anon.get("/cal/wrong-token.ics").status_code == 404
    client.put("/api/calendar/settings", json={"title_mode": "masked", "include_todos": False})
    masked = anon.get(path).text.replace("\r\n ", "")
    assert "髙橋" not in masked and "TODO" not in masked
    assert client.get("/api/calendar/settings").json()["last_access"]

    client.post("/api/calendar/token")          # 再発行で古い URL は無効
    assert anon.get(path).status_code == 404

    did = client.get(f"/api/cases/{cid}").json()["deadlines"][0]["id"]
    one = client.get(f"/api/deadlines/{did}/ics")
    assert one.status_code == 200 and one.text.count("BEGIN:VEVENT") == 1


# ---------------------------------------------------------------- 保管期限

def test_retention(client, loaded, settings):
    kid = loaded[KEIJI]   # 2025-06-30 終結
    r = client.get("/api/retention").json()
    row = next(c for c in r["cases"] if c["id"] == kid)
    assert row["retention_until"] == "2030-06-30" and row["state"] == "保管中" and r["due"] == []
    client.put("/api/settings/retention", json={"years": 1})
    r = client.get("/api/retention").json()
    assert [c["id"] for c in r["due"]] == [kid]
    assert any(n["type"] == "retention" for n in client.get("/api/notifications").json())

    expire_reauth(settings)
    assert client.put(f"/api/cases/{kid}/retention", json={"disposed_on": "2026-10-15"}).status_code == 403
    client.post("/api/auth/reauth", json={"token": TOKEN})
    assert client.put(f"/api/cases/{kid}/retention", json={"disposed_on": "2026-10-15",
                                                           "disposal_note": "溶解処理(業者証明書あり)"}).status_code == 200
    row = next(c for c in client.get("/api/retention").json()["cases"] if c["id"] == kid)
    assert row["state"] == "廃棄済み"
    assert not any(n["type"] == "retention" for n in client.get("/api/notifications").json())
    # 進行中の案件には廃棄を記録できない
    assert client.put(f"/api/cases/{loaded[TAKAHASHI]}/retention", json={"disposed_on": "2026-10-15"}).status_code == 422
    # 個別の保管期限
    client.put(f"/api/cases/{kid}/retention", json={"retention_until": "2040-01-01"})
    assert next(c for c in client.get("/api/retention").json()["cases"] if c["id"] == kid)["retention_until"] == "2040-01-01"


# ---------------------------------------------------------------- 引継ぎ書出し(v3 完了条件)

def test_handover_export(client, loaded, settings, office):
    cid = loaded[TAKAHASHI]
    client.post(f"/api/cases/{cid}/finance", json={"date": "2026-05-15", "kind": "預り金受入", "category": "預り金",
                                                   "amount": 50000})
    client.post(f"/api/cases/{cid}/deadlines", json={"title": "控訴期間", "due_date": "2026-12-01",
                                                     "calc": {"steps": [], "basis": [], "warnings": []}})
    expire_reauth(settings)
    assert client.post("/api/export/handover", json={}).status_code == 403
    client.post("/api/auth/reauth", json={"token": TOKEN})

    r = client.post("/api/export/handover", json={"statuses": ["相談", "受任", "係属", "終結"], "note": "後任: 乙山弁護士"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    text = doc_text(r.content)
    for w in ["事件引継書", "案件一覧(3件)", TAKAHASHI, "斎藤家 遺産分割(被相続人 斎藤 正)", "渡邊 労働事件(解雇無効)",
              "依頼者: 髙橋 一郎", "相手方: 株式会社サンプル商事", "第1回口頭弁論期日", "【計算値・未確定】",
              "相手方答弁書の提出待ち", "預り金残高: 50,000円", "後任: 乙山弁護士", "取扱注意", "テスト法律事務所"]:
        assert w in text, w
    assert "山﨑" not in text                    # 機微案件は既定で除外

    r = client.post("/api/export/handover", json={"case_ids": [loaded[KEIJI]], "include_sensitive": True})
    assert "山﨑 太郎" in doc_text(r.content)
    assert client.post("/api/export/handover", json={"case_ids": [loaded[KEIJI]]}).status_code == 422

    r = client.post("/api/export/handover", json={"format": "csv"})
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("No.,事件名") and len(lines) == 3      # 受任・係属の2件
    log = [a for a in client.get("/api/audit").json() if a["action"] == "引継ぎ書の出力"]
    assert len(log) == 3 and f'"case_ids": [{cid}' in log[-1]["detail"]


# ---------------------------------------------------------------- 受任経路・分野別

def test_analysis(client, loaded):
    t, s = loaded[TAKAHASHI], loaded["斎藤家 遺産分割(被相続人 斎藤 正)"]
    client.put(f"/api/cases/{t}", json={"referral_source": "紹介(弁護士)"})
    client.put(f"/api/cases/{s}", json={"referral_source": "ウェブサイト・広告"})
    client.post(f"/api/cases/{t}/finance", json={"date": "2026-05-15", "kind": "入金", "category": "着手金", "amount": 330000})
    client.post(f"/api/cases/{t}/finance", json={"date": "2026-05-15", "kind": "預り金受入", "category": "預り金", "amount": 99999})
    a = client.get("/api/reports/analysis").json()
    assert a["total_cases"] == 4
    src = {g["key"]: g for g in a["by_source"]}
    assert src["紹介(弁護士)"]["fees_received"] == 330000 and src["紹介(弁護士)"]["hourly"] == 165000
    assert src["未設定"]["cases"] == 2 and src["ウェブサイト・広告"]["minutes"] == 60
    fld = {g["key"]: g for g in a["by_field"]}
    assert fld["民事"]["fees_received"] == 330000 and fld["刑事"]["closed"] == 1 and fld["労働"]["consultations"] == 1
    a = client.get("/api/reports/analysis", params={"date_from": "2026-06-01"}).json()
    assert a["total_cases"] == 2   # 受任日が6月以降の斎藤事件と、受任日なし(登録日=本日)の渡邊事件
